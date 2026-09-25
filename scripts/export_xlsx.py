#!/usr/bin/env python3
"""Export data/leads.db to output/Giraffe_Leads_Master.xlsx and output/Active_Pipeline.csv
(CLAUDE.md Section 11.3).

Sheets, in order: Active_Pipeline, Longterm_Broadcast, Rejected, State_Coverage,
Source_Log, Legend. Active_Pipeline keeps the exact column order of the brief,
sorted Tier, Priority desc, Wave asc, Revenue desc; header row and first four
columns frozen; filters; Tier row colours; clickable links; dropdowns on
Status and FU_Stage. The CSV mirrors Active_Pipeline in UTF-8 with BOM.

Usage:
  python3 scripts/export_xlsx.py [--db data/leads.db] [--out output/Giraffe_Leads_Master.xlsx]
"""
from __future__ import annotations

import argparse
import csv
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, str(Path(__file__).resolve().parent))
import common as C  # noqa: E402
import db  # noqa: E402
import progress  # noqa: E402

TIER_FILL = {"A": "E2F0D9", "B": "FFF2CC", "C": "EDEDED"}          # light green, light amber, light grey
GROUP_FILL = {"Priority": "1F3864", "Identity": "2F5496", "Size": "385723", "Representation": "7B2C2C",
              "Trigger": "B45F06", "People and contact": "5B2C6F", "Checklist": "3A3A3A",
              "Outreach": "0B5345", "Calling tracker": "7F6000"}
SHEET_HEADER_FILL = "2F5496"
URL_COLUMNS = {"Website_URL", "Trigger_Source_URL", "Owner_Mobile_Source_URL", "Next_Gen_LinkedIn_URL",
               "Source_URL", "URL"}
DATE_COLUMNS = {"Trigger_Date", "Date_Researched", "Last_Contact_Date", "Date", "Last_Updated"}
WIDTHS = {
    "Lead_ID": 9, "Tier": 6, "Priority_Score": 9, "Company_Name": 32, "Industry": 22, "Key_Products": 34,
    "State": 14, "City_Cluster": 16, "Plant_Locations": 26, "Wave": 6, "Website_URL": 28, "Year_Established": 9,
    "Legal_Entity": 14, "CIN": 23, "GSTIN": 17, "Revenue_Estimate_Cr": 10, "Revenue_FY": 9, "Revenue_Source": 40,
    "Size_Confidence": 10, "Employees": 12, "Exporter_Markets": 22, "Certifications": 24, "Key_Clients": 28,
    "Website_Decay_Score": 8, "Decay_Evidence": 42, "Corporate_Profile_Status": 26, "LinkedIn_Presence": 26,
    "Social_Presence": 24, "Google_Rating": 14, "Strength_Score": 9, "Visibility_Score": 9, "Gap_Score": 7,
    "Trigger_Type": 18, "Trigger_Detail": 38, "Trigger_Date": 12, "Trigger_Source_URL": 30, "Trigger_Score": 8,
    "Owner_Name": 22, "Owner_Designation": 18, "Owner_Mobile": 17, "Owner_Mobile_Source_URL": 30,
    "Owner_Email": 26, "Next_Gen_Name": 24, "Next_Gen_Role": 18, "Next_Gen_LinkedIn_URL": 30,
    "Company_Phone": 18, "Company_Email": 26, "Email_Type": 12, "Contact_Score": 8, "Manual_Check": 32,
    "Broadcast_Segments": 28, "Recommended_Offer": 22, "Pitch_Hook": 50, "Research_Notes": 50,
    "Date_Researched": 12, "Assigned_To": 14, "FU_Stage": 8, "Last_Contact_Date": 12, "Status": 16,
    "Next_Action": 26, "Drop_Reason": 22, "Reason_For_Longterm": 36, "City": 16, "Reject_Code": 8,
    "Reject_Evidence": 50, "Source_URL": 40, "Date": 12, "Lead_ID_or_Company": 28, "Source_Type": 18,
    "URL": 50, "What_Was_Taken": 50, "Notes": 50, "Last_Updated": 12,
}
WRAP = {"Key_Products", "Plant_Locations", "Revenue_Source", "Exporter_Markets", "Certifications", "Key_Clients",
        "Decay_Evidence", "Corporate_Profile_Status", "LinkedIn_Presence", "Social_Presence", "Trigger_Detail",
        "Manual_Check", "Broadcast_Segments", "Pitch_Hook", "Research_Notes", "Next_Action", "Drop_Reason",
        "Reason_For_Longterm", "Reject_Evidence", "What_Was_Taken", "Notes", "Company_Name"}
THIN = Side(style="thin", color="D9D9D9")


def _link_target(value: str) -> str | None:
    if C.is_blank(value):
        return None
    url = C.first_url(value)
    if url:
        return url
    v = str(value).strip()
    return f"http://{v}" if "." in v and " " not in v and "@" not in v else None


def _cell_value(col: str, value, fill_not_found: bool = True):
    if col in C.NOT_FOUND_EXEMPT or not fill_not_found:
        if value is None or (isinstance(value, str) and not value.strip()):
            return None
    elif value is None or (isinstance(value, str) and not value.strip()):
        return "Not found"
    if col in DATE_COLUMNS and isinstance(value, str):
        d, prec = C.parse_date(value)
        if d and prec == "day":
            return datetime(d.year, d.month, d.day)
    return value


def write_sheet(ws, columns: list[str], rows: list[dict], group_of: dict | None = None,
                freeze: str = "B2", tier_colour: bool = False, fill_not_found: bool = False):
    ws.append(columns)
    for i, col in enumerate(columns, 1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor=(GROUP_FILL.get((group_of or {}).get(col), SHEET_HEADER_FILL)))
        c.alignment = Alignment(wrap_text=True, vertical="center", horizontal="center")
        ws.column_dimensions[get_column_letter(i)].width = WIDTHS.get(col, 14)
    ws.row_dimensions[1].height = 42
    for r_idx, row in enumerate(rows, 2):
        fill = PatternFill("solid", fgColor=TIER_FILL[row["Tier"]]) if tier_colour and row.get("Tier") in TIER_FILL else None
        for c_idx, col in enumerate(columns, 1):
            cell = ws.cell(row=r_idx, column=c_idx, value=_cell_value(col, row.get(col), fill_not_found))
            cell.alignment = Alignment(wrap_text=col in WRAP, vertical="top")
            cell.border = Border(bottom=THIN)
            if fill:
                cell.fill = fill
            if isinstance(cell.value, datetime):
                cell.number_format = "DD-MMM-YYYY"
            target = _link_target(row.get(col)) if (col in URL_COLUMNS or C.first_url(row.get(col))) else None
            if target:
                cell.hyperlink = target
                cell.font = Font(color="1F4E79", underline="single" if col in URL_COLUMNS else None)
    ws.freeze_panes = freeze
    last_row = max(len(rows) + 1, 2)
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last_row}"


def add_dropdowns(ws, columns: list[str], n_rows: int):
    last = max(n_rows + 1, 2000)
    for col, values in (("Status", C.STATUS_VALUES), ("FU_Stage", C.FU_STAGES)):
        letter = get_column_letter(columns.index(col) + 1)
        dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True,
                            showErrorMessage=True, errorTitle=f"Invalid {col}",
                            error=f"{col} must be one of: {', '.join(values)}")
        dv.add(f"{letter}2:{letter}{last}")
        ws.add_data_validation(dv)


def legend_rows() -> list[tuple]:
    rows = [("GIRAFFE INDUSTRIAL MARKETING: LEAD MASTER LEGEND", "", "", ""),
            (f"Generated {C.today_str()} from data/leads.db by scripts/export_xlsx.py", "", "", ""),
            ("", "", "", ""), ("COLUMNS", "Column", "Group", "Definition")]
    rows += [("Active_Pipeline", n, g, d) for n, _, g, d in C.ACTIVE_SPECS]
    rows += [("Longterm_Broadcast", "Lead_ID + Identity + People and contact columns", "", "Same definitions as Active_Pipeline.")]
    rows += [("Longterm_Broadcast", n, g, d) for n, _, g, d in C.LONGTERM_EXTRA_SPECS]
    rows += [("Rejected", n, "", d) for n, _, _, d in C.REJECTED_SPECS]
    rows += [("State_Coverage", n, "", d) for n, _, _, d in C.COVERAGE_SPECS]
    rows += [("Source_Log", n, "", d) for n, _, _, d in C.SOURCE_LOG_SPECS]
    rows += [("", "", "", ""), ("SCORING (Section 10)", "Rule", "", "Points")]
    rows += [
        ("Strength_Score", "Revenue", "", "₹50-100 Cr = 1, ₹100-250 Cr = 2, ₹250-500 Cr = 3"),
        ("Strength_Score", "Proof points", "", "Certifications, exports, OEM or institutional clients, awards: 0 to 3"),
        ("Strength_Score", "Scale", "", "Plant capacity, multiple plants, 200+ employees: 0 to 2"),
        ("Strength_Score", "Longevity", "", "15+ years in business = 1"),
        ("Strength_Score", "Reach", "", "Dealers in 3+ states or exports to 3+ countries = 1"),
        ("Visibility_Score", "Website", "", "Decay 0-2 = 4, 3-4 = 3, 5-6 = 2, 7-8 = 1, 9-10 = 0"),
        ("Visibility_Score", "Corporate profile", "", "Current = 2, Old = 1, None = 0"),
        ("Visibility_Score", "LinkedIn page", "", "Posts in last 90 days = 2, inactive = 1, none = 0"),
        ("Visibility_Score", "Google Business", "", "Rating 4+ with 20+ reviews = 1"),
        ("Visibility_Score", "Leadership visible online", "", "Interviews or active LinkedIn = 1"),
        ("Gap_Score", "Strength - Visibility", "", "-10 to +10: the representation gap"),
        ("Trigger_Score", "Highest single trigger", "", "0 to 5, within 18 months (trade fairs: 2025 or 2026 editions)"),
        ("Contact_Score", "Best contact found", "", "3 owner/director mobile with source URL; 2 owner direct email or next-gen LinkedIn; 1 company line or generic email; 0 none"),
        ("Priority_Score", "Formula", "", "(Gap x 2) + (Trigger x 3) + (Contact x 2) + Wave bonus (Wave 1 +3, Wave 2 +1)"),
        ("Tier A", "", "", "Gap 5+ AND Trigger 3+ AND Size_Confidence High or Medium AND Contact 2+"),
        ("Tier B", "", "", "Gap 3+ AND (Trigger 2+ OR Contact 2+)"),
        ("Tier C", "", "", "Passes all gates but weaker, or Size_Confidence Low (with Manual_Check)"),
        ("Recommended_Offer", "Rule order", "", "CAPMKT or Pvt to Public -> PREMIUM + CAPITAL MARKETS; else exhibitor (SEG-EXPO) -> EXPO KIT; else revenue ₹150 Cr+ -> PREMIUM + CAPITAL MARKETS; else ESSENTIALS + RETAINER. A researcher override is allowed and flagged."),
        ("Automatic segments", "", "", "Industry segment; SEG-EXPO for a trade fair trigger; SEG-EXPORT when export markets are named; SEG-NEXTGEN when a next-gen director is named; SEG-NOWEB when there is no website."),
        ("Row colours", "", "", "Tier A light green, Tier B light amber, Tier C light grey"),
    ]
    rows += [("", "", "", ""), ("REJECT CODES", "Code", "", "Meaning")]
    rows += [("Reject code", k, "", v) for k, v in C.REJECT_CODES.items()]
    rows += [("", "", "", ""), ("BROADCAST SEGMENTS", "Tag", "", "Meaning")]
    rows += [("Segment", k, "", v) for k, v in C.SEGMENTS.items()]
    rows += [("Segment", "SRC-*", "", "Source list tag for rows from input/ lists, e.g. SRC-FRANCHISE-INDIA")]
    rows += [("", "", "", ""), ("TRIGGER TYPES", "Type", "Score", "Meaning")]
    rows += [("Trigger type", k, str(s), d) for k, (s, d) in C.TRIGGER_TYPES.items()]
    rows += [("", "", "", ""), ("RECOMMENDED OFFERS", "Offer", "", "When")]
    rows += [("Offer", k, "", v) for k, v in C.OFFERS.items()]
    rows += [("", "", "", ""), ("OTHER CODES", "Field", "", "Allowed values")]
    rows += [("Industry", "", "", ", ".join(C.INDUSTRIES)),
             ("Status", "", "", ", ".join(C.STATUS_VALUES)),
             ("FU_Stage", "", "", "0, 1, 2, 3 (maximum 3 follow-ups, then drop)"),
             ("Size_Confidence", "", "", "High: rating rationale, audited financials, DRHP or company figure with FY. "
                                         "Medium: IndiaMART/TradeIndia band, credible news, GST slab. Low: proxies only."),
             ("Checklist columns", "", "", "Y or N. Manual_Check is blank or a note of what to verify by hand."),
             ("State_Coverage Status", "", "", ", ".join(C.COVERAGE_STATUSES)),
             ("Unknown values", "", "", "'Not found' means searched and not found; '(inferred)' marks an inferred value; "
                                        "'(probable)' marks an unconfirmed next-gen relationship.")]
    return rows


def export(db_path: Path = C.DB_PATH, out_xlsx: Path = C.OUTPUT_DIR / "Giraffe_Leads_Master.xlsx",
           out_csv: Path = C.OUTPUT_DIR / "Active_Pipeline.csv") -> dict:
    conn = db.connect(db_path)
    active = [dict(r) for r in conn.execute(
        "SELECT * FROM leads WHERE Pipeline = 'ACTIVE' ORDER BY CASE Tier WHEN 'A' THEN 0 WHEN 'B' THEN 1 ELSE 2 END,"
        " Priority_Score DESC, Wave ASC, Revenue_Estimate_Cr DESC")]
    longterm = [dict(r) for r in conn.execute("SELECT * FROM leads WHERE Pipeline = 'LONGTERM' ORDER BY Wave, State, Company_Name")]
    rejected = [dict(r) for r in conn.execute("SELECT * FROM rejected ORDER BY State, City, Company_Name")]
    sources = [dict(r) for r in conn.execute("SELECT * FROM source_log ORDER BY id")]
    cov_rows = progress.coverage_rows(progress.load_coverage(), db.counts_by_state_entry(conn))

    wb = Workbook()
    ws = wb.active
    ws.title = "Active_Pipeline"
    group_of = {n: g for n, _, g, _ in C.ACTIVE_SPECS}
    write_sheet(ws, C.ACTIVE_COLUMNS, active, group_of, freeze="E2", tier_colour=True, fill_not_found=True)
    add_dropdowns(ws, C.ACTIVE_COLUMNS, len(active))

    group_lt = dict(group_of, Reason_For_Longterm="Outreach")
    write_sheet(wb.create_sheet("Longterm_Broadcast"), C.LONGTERM_COLUMNS, longterm, group_lt, freeze="C2",
                fill_not_found=True)
    write_sheet(wb.create_sheet("Rejected"), C.REJECTED_COLUMNS, rejected, freeze="B2")
    ws_cov = wb.create_sheet("State_Coverage")
    write_sheet(ws_cov, C.COVERAGE_COLUMNS, cov_rows, freeze="C2")
    status_col = C.COVERAGE_COLUMNS.index("Status") + 1
    for r in range(2, len(cov_rows) + 2):
        cell = ws_cov.cell(row=r, column=status_col)
        colour = {"Done": "E2F0D9", "In Progress": "FFF2CC", "N/A": "EDEDED"}.get(cell.value)
        if colour:
            cell.fill = PatternFill("solid", fgColor=colour)
    write_sheet(wb.create_sheet("Source_Log"), C.SOURCE_LOG_COLUMNS, sources, freeze="B2")

    lg = wb.create_sheet("Legend")
    for row in legend_rows():
        lg.append(list(row))
    for i, w in enumerate((26, 34, 18, 110), 1):
        lg.column_dimensions[get_column_letter(i)].width = w
    for row in lg.iter_rows():
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        first = row[0].value or ""
        if first.isupper() and first:
            for cell in row:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor=SHEET_HEADER_FILL)
    lg.freeze_panes = "A5"

    out_xlsx.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_xlsx)
    with open(out_csv, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow(C.ACTIVE_COLUMNS)
        for row in active:
            vals = []
            for col in C.ACTIVE_COLUMNS:
                v = _cell_value(col, row.get(col))
                vals.append(C.fmt_date(v.date()) if isinstance(v, datetime) else ("" if v is None else v))
            w.writerow(vals)
    return {"active": len(active), "longterm": len(longterm), "rejected": len(rejected), "sources": len(sources),
            "xlsx": str(out_xlsx), "csv": str(out_csv)}


def main(argv=None) -> int:
    p = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--db", default=str(C.DB_PATH))
    p.add_argument("--out", default=str(C.OUTPUT_DIR / "Giraffe_Leads_Master.xlsx"))
    p.add_argument("--csv", default=str(C.OUTPUT_DIR / "Active_Pipeline.csv"))
    args = p.parse_args(argv)
    res = export(Path(args.db), Path(args.out), Path(args.csv))
    print(f"Exported {res['active']} active, {res['longterm']} longterm, {res['rejected']} rejected, "
          f"{res['sources']} sources -> {res['xlsx']} and {res['csv']}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
