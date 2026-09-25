"""candidates_xlsx.py: merge, screening and workbook on fictional research output."""
import json
import sys
from datetime import date
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
import candidates_xlsx as X  # noqa: E402

TODAY = date(2026, 9, 25)
A = {"candidates": [
    {"company_name": "Fictional Knit Pvt Ltd", "is_manufacturer": "Yes", "industry": "Apparel and Hosiery",
     "revenue_cr": 180, "revenue_fy": "FY2025", "revenue_source": "Acuite 12-Jan-2026",
     "revenue_url": "https://www.acuite.in/x.pdf", "expansion_signal": "Debt-funded capex for new knitting unit",
     "expansion_date": "12-Jan-2026", "expansion_url": "https://www.acuite.in/x.pdf",
     "popularity_or_group_flags": "None found", "website": "fictionalknit.example"},
    {"company_name": "Fictional Listed Forge Ltd", "is_manufacturer": "Yes", "revenue_cr": 300, "revenue_fy": "FY2025",
     "popularity_or_group_flags": "Listed on NSE main board", "flags_url": "https://nse.example"},
    {"company_name": "Fictional Big Steel Pvt Ltd", "is_manufacturer": "Yes", "revenue_cr": 900, "revenue_fy": "FY2025",
     "revenue_url": "https://careratings.com/y.pdf"},
    {"company_name": "Fictional Traders", "is_manufacturer": "No", "manufacturer_evidence": "trading house"},
], "screened_out": [{"company_name": "Fictional Hotels Pvt Ltd", "reason": "hotel", "url": "https://h.example"}]}
B = {"candidates": [
    {"company_name": "FICTIONAL KNIT PRIVATE LIMITED", "trade_fair": "Knit Show 2026", "trade_fair_date": "Mar-2026",
     "trade_fair_url": "https://fair.example/2026", "promoters": "A. Fictional"},
    {"company_name": "Fictional Cycle Parts", "is_manufacturer": "Yes", "revenue_cr": 75, "revenue_fy": "FY2025",
     "revenue_source": "IndiaMART band Rs 50-100 Crore", "revenue_url": "https://www.indiamart.com/fcp/",
     "expansion_signal": "New plant", "expansion_date": "July 17, 2024", "expansion_url": "https://news.example/a",
     "popularity_or_group_flags": "None found"},
    {"company_name": "Fictional Tools", "is_manufacturer": "Unclear"},
], "screened_out": []}


def test_merge_screen_and_workbook(tmp_path):
    (tmp_path / "a.json").write_text(json.dumps(A))
    (tmp_path / "b.json").write_text(json.dumps(B))
    res = X.build([tmp_path / "a.json", tmp_path / "b.json"], "Ludhiana", "Punjab", "LDH",
                  tmp_path / "c.xlsx", use_db=False, today=TODAY)
    assert (res["candidates"], res["Strong"], res["Good"], res["Check"]) == (2, 1, 1, 0)
    wb = load_workbook(tmp_path / "c.xlsx")
    assert wb.sheetnames == ["Read_Me", "Candidates", "Follow_Up", "Screened_Out"]
    assert [r[0].value for r in wb["Follow_Up"].iter_rows(min_row=2)] == ["Fictional Tools"]
    ws = wb["Candidates"]
    head = [c.value for c in ws[1]]
    first = {h: ws.cell(row=2, column=i + 1).value for i, h in enumerate(head)}
    assert first["Cand_ID"] == "LDH-C01" and first["Fit"] == "Strong"
    assert first["Company_Name"] == "Fictional Knit Pvt Ltd"          # merged across both files
    assert first["Signal_Type"] == "Rating capex" and first["Promoters"] == "A. Fictional"
    assert "Trade fair" in first["Other_Signals"]
    second = {h: ws.cell(row=3, column=i + 1).value for i, h in enumerate(head)}
    assert second["Fit"] == "Good" and second["Signal_Type"] == "None within window"   # 2024 news is too old
    assert "self-declared" in second["Size_Confidence"]
    reasons = {r[0].value: r[1].value for r in wb["Screened_Out"].iter_rows(min_row=2)}
    assert reasons == {"Fictional Listed Forge Ltd": "R5", "Fictional Big Steel Pvt Ltd": "R1",
                       "Fictional Traders": "R7", "Fictional Hotels Pvt Ltd": "R7"}


def test_review_note_keeps_flagged_company_visible(tmp_path):
    data = {"candidates": [{"company_name": "Fictional Springs Ltd", "is_manufacturer": "Yes", "revenue_cr": 114,
                            "revenue_fy": "FY2025", "popularity_or_group_flags": "A venture capital firm holds 20%",
                            "review_note": "Legacy stake: decide R4"},
                           {"company_name": "Fictional Tyres", "is_manufacturer": "Yes", "revenue_cr": 325,
                            "revenue_fy": "FY2023",
                            "popularity_or_group_flags": "No stock listing, large-group link or Wikipedia page found"}],
            "screened_out": []}
    (tmp_path / "r.json").write_text(json.dumps(data))
    res = X.build([tmp_path / "r.json"], "Ludhiana", "Punjab", "LDH", tmp_path / "r.xlsx", use_db=False, today=TODAY)
    assert (res["candidates"], res["Good"], res["Check"]) == (2, 1, 1)
    ws = load_workbook(tmp_path / "r.xlsx")["Candidates"]
    head = [c.value for c in ws[1]]
    rows = [{h: ws.cell(row=r, column=i + 1).value for i, h in enumerate(head)} for r in (2, 3)]
    assert rows[0]["Company_Name"] == "Fictional Tyres" and rows[0]["Fit"] == "Good"
    assert rows[1]["Fit"] == "Check" and "REVIEW: Legacy stake" in rows[1]["Screen_Check"]


def test_enrichment_adds_site_directors_and_ranked_contact_pages(tmp_path):
    data = {"candidates": [{"company_name": "Fictional Gears Ltd", "is_manufacturer": "Yes", "revenue_cr": 120,
                            "revenue_fy": "FY2025", "popularity_or_group_flags": "None found"}], "screened_out": []}
    enrich = [{"company_name": "Fictional Gears Limited", "website": "fictionalgears.example",
               "directors": [{"name": "A. Fictional", "designation": "Managing Director", "appointed": "1995"},
                             {"name": "B. Fictional", "designation": "Director", "appointed": "2021",
                              "note": "next-gen (probable)"}],
               "owner_contact_pages": [
                   {"type": "Company contact page", "url": "https://fictionalgears.example/contact"},
                   {"type": "Other", "url": "https://rocketreach.co/fictional-gears"},
                   {"type": "IndiaMART", "url": "https://www.indiamart.com/fictional-gears/",
                    "listed_contact_person": "A. Fictional (MD)"}]}]
    (tmp_path / "c.json").write_text(json.dumps(data))
    (tmp_path / "e.json").write_text(json.dumps(enrich))
    X.build([tmp_path / "c.json"], "Ludhiana", "Punjab", "LDH", tmp_path / "e.xlsx", use_db=False, today=TODAY,
            enrich=[tmp_path / "e.json"])
    ws = load_workbook(tmp_path / "e.xlsx")["Candidates"]
    head = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(head)}
    assert ws.cell(row=2, column=col["Website_URL"]).value == "fictionalgears.example"
    assert "B. Fictional (Director), since 2021" == ws.cell(row=2, column=col["Next_Gen"]).value
    first = ws.cell(row=2, column=col["Contact_Page_1"])
    assert first.value == "IndiaMART: A. Fictional (MD)" and first.hyperlink.target == "https://www.indiamart.com/fictional-gears/"
    assert ws.cell(row=2, column=col["Contact_Page_2"]).value == "Company contact page"
    assert not ws.cell(row=2, column=col["Contact_Page_3"]).value          # the data-broker link was dropped
    assert ws.cell(row=2, column=col["Owner_Mobile"]).value.startswith(
        "Likely on Contact_Page_1 (IndiaMART), which names A. Fictional (MD)")


def test_doubtful_director_label_is_not_treated_as_owner(tmp_path):
    data = {"candidates": [{"company_name": "Fictional Looms Ltd", "is_manufacturer": "Yes", "revenue_cr": 90,
                            "revenue_fy": "FY2025", "popularity_or_group_flags": "None found"}], "screened_out": []}
    enrich = [{"company_name": "Fictional Looms Ltd", "website": "Not found",
               "directors": [{"name": "C. Fictional", "designation": "Director", "note": "not next-gen: status unverified"}],
               "owner_contact_pages": [{"type": "TradeIndia", "url": "https://www.tradeindia.com/fictional-looms/",
                                        "listed_contact_person": "Mr. X, Director (not an MCA director, likely staff)"}]}]
    (tmp_path / "c.json").write_text(json.dumps(data))
    (tmp_path / "e.json").write_text(json.dumps(enrich))
    X.build([tmp_path / "c.json"], "Ludhiana", "Punjab", "LDH", tmp_path / "d.xlsx", use_db=False, today=TODAY,
            enrich=[tmp_path / "e.json"])
    ws = load_workbook(tmp_path / "d.xlsx")["Candidates"]
    col = {c.value: i + 1 for i, c in enumerate(ws[1])}
    assert ws.cell(row=2, column=col["Owner_Mobile"]).value.startswith("Open Contact_Page_1-3")
    assert ws.cell(row=2, column=col["Next_Gen"]).value == "Not found"
