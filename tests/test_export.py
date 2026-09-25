"""export_xlsx.py and progress.py on a temporary DB of fictional leads."""
import csv
import sys
from pathlib import Path

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "scripts"))
sys.path.insert(0, str(Path(__file__).resolve().parent))
import common as C  # noqa: E402
import db  # noqa: E402
import export_xlsx  # noqa: E402
import progress  # noqa: E402
from test_scoring_db import ev, lead  # noqa: E402


def build_db(path):
    conn = db.connect(path)
    for d in (lead(Company_Name="Fictional B Tools", Website_URL="https://fictional-b.example", GSTIN=None),
              lead(Company_Name="Fictional A Forge", Website_URL="https://fictional-a.example", GSTIN=None,
                   Website_Decay_Score=9),
              lead(Company_Name="Fictional C Mills", Website_URL="https://fictional-c.example", GSTIN=None,
                   Size_Confidence="Low", Manual_Check="Verify turnover", Trigger_Date="Aug-2026"),
              lead(Pipeline="LONGTERM", Company_Name="Fictional Roads EPC", Website_URL="https://fictional-r.example",
                   GSTIN=None, Industry="Building Materials", Broadcast_Segments="SEG-GIFTING",
                   Reason_For_Longterm="Road EPC contractor")):
        r = ev(d)
        assert not r.errors, r.errors
        db.upsert_lead(conn, r.row)
    db.add_reject(conn, {"Company_Name": "Fictional Listed Ltd", "State": "Punjab", "City": "Ludhiana",
                         "Reject_Code": "R5", "Reject_Evidence": "NSE main board", "Source_URL": "https://x.example"})
    db.log_source(conn, "PB-001", "Rating rationale", "https://ratings.example.org/r/1.pdf", "FY25 revenue",
                  "Punjab", "Ludhiana")
    return conn


def test_workbook_structure_and_csv(tmp_path):
    build_db(tmp_path / "t.db")
    res = export_xlsx.export(tmp_path / "t.db", tmp_path / "m.xlsx", tmp_path / "a.csv")
    assert (res["active"], res["longterm"], res["rejected"]) == (3, 1, 1)
    wb = load_workbook(tmp_path / "m.xlsx")
    assert wb.sheetnames == ["Active_Pipeline", "Longterm_Broadcast", "Rejected", "State_Coverage",
                             "Source_Log", "Legend"]
    ws = wb["Active_Pipeline"]
    header = [c.value for c in ws[1]]
    assert header == C.ACTIVE_COLUMNS
    assert ws.freeze_panes == "E2" and ws.auto_filter.ref.startswith("A1:")
    tiers = [ws.cell(row=r, column=2).value for r in range(2, 5)]
    assert tiers == ["A", "B", "C"]                      # sorted by tier
    assert ws.cell(row=2, column=1).fill.fgColor.rgb.endswith("E2F0D9")   # Tier A light green
    web_col = header.index("Website_URL") + 1
    assert ws.cell(row=2, column=web_col).hyperlink.target == "https://fictional-a.example"
    dv_cols = {str(dv.sqref).split(":")[0].rstrip("0123456789") for dv in ws.data_validations.dataValidation}
    letters = {ws.cell(row=1, column=header.index(c) + 1).column_letter for c in ("Status", "FU_Stage")}
    assert dv_cols == letters
    assert ws.cell(row=2, column=header.index("GSTIN") + 1).value == "Not found"
    assert ws.cell(row=2, column=header.index("Assigned_To") + 1).value is None
    assert ws.cell(row=2, column=header.index("Date_Researched") + 1).number_format == "DD-MMM-YYYY"
    assert [c.value for c in wb["Rejected"][1]] == C.REJECTED_COLUMNS
    assert [c.value for c in wb["State_Coverage"][1]] == C.COVERAGE_COLUMNS
    cov = {r[1].value: r for r in wb["State_Coverage"].iter_rows(min_row=2)}
    assert cov["Punjab"][C.COVERAGE_COLUMNS.index("Candidates_Screened")].value == 5
    assert len(cov) == len(C.state_entries())
    raw = (tmp_path / "a.csv").read_bytes()
    assert raw.startswith(b"\xef\xbb\xbf")               # UTF-8 BOM for Excel
    rows = list(csv.reader(raw.decode("utf-8-sig").splitlines()))
    assert rows[0] == C.ACTIVE_COLUMNS and len(rows) == 4
    assert rows[1][header.index("Date_Researched")] == "25-Sep-2026"


def test_checklist_render_keeps_manual_header(tmp_path):
    cov = {"states": {}}
    progress.mark(cov, "Punjab", "Ludhiana", ["ratings", "ec_cte"], False, "stop rule not yet hit")
    progress.mark(cov, "UP", "Meerut", ["expos"], True, None)
    rows = progress.coverage_rows(cov, {"Punjab": {"A": 1, "B": 2, "C": 0, "longterm": 0, "rejected": 4,
                                                   "sources": 9, "screened": 7}})
    path = tmp_path / "STATE_CHECKLIST.md"
    path.write_text("# My header\nmanual notes\n")
    progress.write_checklist(rows, path)
    progress.write_checklist(rows, path)                # idempotent
    text = path.read_text()
    assert text.startswith("# My header\nmanual notes\n") and text.count(progress.BEGIN) == 1
    assert "### Punjab [IN PROGRESS] screened: 7 | A: 1 | B: 2 | C: 0 | longterm: 0 | rejected: 4" in text
    assert "- [ ] Ludhiana: [x] ratings [x] EC/CTE [ ] expos" in text
    assert "### Uttar Pradesh (West and NCR belt) [IN PROGRESS]" in text
    assert "- [x] Meerut: [ ] ratings [ ] EC/CTE [x] expos" in text
